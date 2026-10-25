"""
ingest.py — reads OLGA multiphase-flow export workbooks.

Why this file exists: the workbooks store every number INSIDE an Excel formula
(for example  =E$9 * (0.0189 - E$8) ), and the cached results are not reliable.
A plain pandas.read_excel() therefore returns zeros or blanks. This reader
rebuilds each value from the formula text and identifies the columns from the
workbook's own 'Legends' sheet. Same code as the Colab notebook pipeline.
"""
import re
import glob
import os
import openpyxl
import pandas as pd
import numpy as np

VALUE_PATTERN = re.compile(r'\(\s*([\-0-9.eE]+)\s*-')


def extract_value(cell_value):
    """Real number from a plain number or an Offset/Scale formula string."""
    if cell_value is None:
        return np.nan
    if isinstance(cell_value, (int, float)):
        return float(cell_value)
    if isinstance(cell_value, str):
        m = VALUE_PATTERN.search(cell_value)
        if m:
            return float(m.group(1))
        try:
            return float(cell_value)
        except ValueError:
            return np.nan
    return np.nan


def _read_legend(ws_legend):
    legend = {}
    for i, row in enumerate(ws_legend.iter_rows(values_only=True), start=1):
        if row and row[0] is not None:
            legend[i] = (row[0], row[1] if len(row) > 1 else None)
    return legend


def _friendly_name(description):
    """'TREND___QLT___(M3_S)___RISER_OUTLET' -> ('QLT', 'RISER_OUTLET')."""
    parts = str(description).split('___')
    varname = parts[1] if len(parts) > 1 else str(description)
    location = parts[3] if len(parts) > 3 else None
    return varname, location


def load_workbook_scenarios(path):
    """Every data sheet of one workbook -> list of
    {'genkey', 'source_file', 'source_sheet', 'df'} (df has TIME_h + variables)."""
    wb = openpyxl.load_workbook(path, data_only=False)
    if 'Legends' not in wb.sheetnames:
        return []
    legend = _read_legend(wb['Legends'])

    results = []
    for sheet_name in wb.sheetnames:
        if sheet_name == 'Legends':
            continue
        ws = wb[sheet_name]
        max_col, max_row = ws.max_column, ws.max_row

        time_cols, data_cols, genkey = [], {}, None
        for c in range(1, max_col + 1):
            header = ws.cell(row=2, column=c).value
            if header is None:
                continue
            if isinstance(header, str) and header.startswith('TIME'):
                time_cols.append(c)
            elif isinstance(header, str) and '.genkey' in header:
                genkey = header
            elif isinstance(header, str) and 'Legends' in header:
                m = re.search(r'B(\d+)', header)
                if m and int(m.group(1)) in legend:
                    data_cols[c] = _friendly_name(legend[int(m.group(1))][0])

        if genkey is None or not data_cols or not time_cols:
            continue

        name_counts = {}
        for c, (varname, location) in data_cols.items():
            name_counts[varname] = name_counts.get(varname, 0) + 1
        col_names = {}
        for c, (varname, location) in data_cols.items():
            col_names[c] = f"{varname}__{location}" if (name_counts[varname] > 1 and location) else varname

        # Real data starts after the LAST literal 'TIME1' header row (OLGA repeats
        # the header once, after a block of MIN/MEDIAN/MAX and Offset/Scale rows).
        first_time_col = time_cols[0]
        header_rows = [r for r in range(1, max_row + 1)
                       if ws.cell(row=r, column=first_time_col).value == 'TIME1']
        if not header_rows:
            continue
        start = header_rows[-1] + 1

        data = {'TIME_h': [extract_value(ws.cell(row=r, column=first_time_col).value)
                           for r in range(start, max_row + 1)]}
        for c, name in col_names.items():
            data[name] = [extract_value(ws.cell(row=r, column=c).value)
                          for r in range(start, max_row + 1)]
        df = pd.DataFrame(data).dropna(how='all').reset_index(drop=True)
        results.append({'genkey': genkey, 'source_file': os.path.basename(path),
                        'source_sheet': sheet_name, 'df': df})
    return results


def load_all_workbooks(data_dir, pattern="*.xlsx"):
    all_scenarios = []
    for path in sorted(glob.glob(os.path.join(data_dir, pattern))):
        try:
            all_scenarios.extend(load_workbook_scenarios(path))
        except Exception as e:  # keep going if one file is unreadable
            print(f"[ingest] WARNING: could not read {path}: {e}")
    return all_scenarios
